from __future__ import annotations

import sqlite3
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..db import get_settings
from ..utils import (
    SMS_STATUS_AR,
    STATUS_AR,
    STATUS_ORDER,
    month_bounds,
    month_label,
    short_date,
    validate_iso_date,
)


def _class_sql(class_group: str | None) -> tuple[str, list[Any]]:
    if class_group == "__none__":
        return " AND s.class_group = ''", []
    if class_group:
        return " AND s.class_group = ?", [class_group]
    return "", []


def daily_report(
    conn: sqlite3.Connection,
    day_iso: str,
    period_id: int | None = None,
    class_group: str | None = None,
    status: str | None = None,
) -> dict[str, Any]:
    day_iso = validate_iso_date(day_iso)
    filters = ["a.attendance_date = ?"]
    params: list[Any] = [day_iso]
    if period_id:
        filters.append("a.period_id = ?")
        params.append(int(period_id))
    class_clause, class_params = _class_sql(class_group)
    filters.append("1=1" + class_clause)
    params.extend(class_params)
    if status in STATUS_ORDER:
        filters.append("a.status = ?")
        params.append(status)
    where = " AND ".join(filters)
    rows = conn.execute(
        f"""SELECT a.*, s.first_name, s.last_name, s.class_group,
                   p.period_name, p.start_time,
                   sm.delivery_status AS sms_status, sm.is_test AS sms_is_test
            FROM attendance_logs a
            JOIN students s ON s.student_nin = a.student_nin
            JOIN periods p ON p.id = a.period_id
            LEFT JOIN sms_logs sm ON sm.student_nin = a.student_nin
                 AND sm.period_id = a.period_id AND sm.attendance_date = a.attendance_date
            WHERE {where}
            ORDER BY p.sort_order, p.start_time, CASE WHEN s.class_group = '' THEN 1 ELSE 0 END,
                     s.class_group, s.last_name, s.first_name""",
        params,
    ).fetchall()
    items = []
    summary = {status_key: 0 for status_key in STATUS_ORDER}
    for row in rows:
        data = dict(row)
        summary[data["status"]] = summary.get(data["status"], 0) + 1
        items.append(
            {
                "student_nin": data["student_nin"],
                "first_name": data["first_name"],
                "last_name": data["last_name"],
                "full_name": f"{data['first_name']} {data['last_name']}".strip(),
                "class_group": data["class_group"] or "بدون قسم",
                "period_name": data["period_name"],
                "scan_time": (data["scan_time"] or "")[:5],
                "status": data["status"],
                "status_label": STATUS_AR.get(data["status"], data["status"]),
                "sms_status": data.get("sms_status"),
                "sms_status_label": SMS_STATUS_AR.get(data.get("sms_status") or "", "—"),
                "sms_is_test": bool(data.get("sms_is_test")),
                "note": data["note"] or "",
            }
        )
    settings = get_settings(conn)
    period_label = "كل الفترات"
    if period_id:
        period = conn.execute("SELECT period_name FROM periods WHERE id = ?", (period_id,)).fetchone()
        period_label = period["period_name"] if period else "—"
    class_label = "كل الأقسام"
    if class_group == "__none__":
        class_label = "بدون قسم"
    elif class_group:
        class_label = class_group
    return {
        "kind": "daily",
        "date": day_iso,
        "date_display": short_date(day_iso),
        "period_label": period_label,
        "class_label": class_label,
        "status_filter": status if status in STATUS_ORDER else "",
        "items": items,
        "summary": summary,
        "total": len(items),
        "school_name": settings.get("school_name", ""),
        "academic_year": settings.get("academic_year", ""),
        "header_text": settings.get("header_text", ""),
        "principal_name": settings.get("principal_name", ""),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "title": "تقرير الحضور والغياب اليومي",
    }


def monthly_report(
    conn: sqlite3.Connection,
    month: str,
    period_id: int | None = None,
    class_group: str | None = None,
) -> dict[str, Any]:
    start, end = month_bounds(month)
    start_iso, end_iso = start.isoformat(), end.isoformat()
    period_clause = " AND a.period_id = ?" if period_id else ""
    class_clause, class_params = _class_sql(class_group)
    params: list[Any] = [start_iso, end_iso]
    if period_id:
        params.append(int(period_id))
    params.extend(class_params)
    rows = conn.execute(
        f"""SELECT s.student_nin, s.first_name, s.last_name, s.class_group,
                   SUM(CASE WHEN a.status = 'PRESENT' THEN 1 ELSE 0 END) AS present,
                   SUM(CASE WHEN a.status = 'LATE' THEN 1 ELSE 0 END) AS late,
                   SUM(CASE WHEN a.status = 'ABSENT' THEN 1 ELSE 0 END) AS absent,
                   SUM(CASE WHEN a.status = 'EXCUSED' THEN 1 ELSE 0 END) AS excused,
                   COUNT(a.id) AS total
            FROM students s
            LEFT JOIN attendance_logs a ON a.student_nin = s.student_nin
                 AND a.attendance_date BETWEEN ? AND ?{period_clause}
            WHERE (s.is_active = 1 OR a.id IS NOT NULL){class_clause}
            GROUP BY s.student_nin
            ORDER BY CASE WHEN s.class_group = '' THEN 1 ELSE 0 END, s.class_group,
                     s.last_name, s.first_name""",
        params,
    ).fetchall()
    items = []
    totals = {"present": 0, "late": 0, "absent": 0, "excused": 0, "total": 0}
    for row in rows:
        data = dict(row)
        data["full_name"] = f"{data['first_name']} {data['last_name']}".strip()
        data["class_group"] = data["class_group"] or "بدون قسم"
        for key in totals:
            totals[key] += int(data[key] or 0)
        items.append(data)
    settings = get_settings(conn)
    period_label = "كل الفترات"
    if period_id:
        period = conn.execute("SELECT period_name FROM periods WHERE id = ?", (period_id,)).fetchone()
        period_label = period["period_name"] if period else "—"
    class_label = "كل الأقسام"
    if class_group == "__none__":
        class_label = "بدون قسم"
    elif class_group:
        class_label = class_group
    return {
        "kind": "monthly",
        "month": month,
        "month_display": month_label(month),
        "start": start_iso,
        "end": end_iso,
        "period_label": period_label,
        "class_label": class_label,
        "items": items,
        "totals": totals,
        "total_students": len(items),
        "school_name": settings.get("school_name", ""),
        "academic_year": settings.get("academic_year", ""),
        "header_text": settings.get("header_text", ""),
        "principal_name": settings.get("principal_name", ""),
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "title": "تقرير الغياب الشهري",
    }


def _style_workbook(ws, title: str, report: dict[str, Any], columns: list[tuple[str, int]], header_row: int) -> None:
    ws.sheet_view.rightToLeft = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_options.horizontalCentered = True

    title_font = Font(name="Arial", bold=True, size=16, color="0B3B31")
    header_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    info_font = Font(name="Arial", size=10, color="536D66")
    accent_fill = PatternFill("solid", fgColor="0E7C66")
    light_fill = PatternFill("solid", fgColor="E7F3EE")
    thin = Side(style="thin", color="C9D8D3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    width = len(columns)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(row=1, column=1, value=report.get("school_name") or "المؤسسة التعليمية")
    cell.font = title_font
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    cell = ws.cell(row=2, column=1, value=title)
    cell.font = Font(name="Arial", bold=True, size=13, color="0B3B31")
    cell.alignment = Alignment(horizontal="center", vertical="center")
    if report["kind"] == "daily":
        details = (
            f"التاريخ: {report['date_display']}    |    الفترة: {report['period_label']}    |    "
            f"القسم: {report['class_label']}    |    السنة الدراسية: {report['academic_year']}"
        )
    else:
        details = (
            f"الشهر: {report['month_display']}    |    الفترة: {report['period_label']}    |    "
            f"القسم: {report['class_label']}    |    السنة الدراسية: {report['academic_year']}"
        )
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=width)
    cell = ws.cell(row=3, column=1, value=details)
    cell.font = info_font
    cell.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=width)
    cell = ws.cell(row=4, column=1, value=f"أُنشئ التقرير بواسطة نظام متابعة الحضور — {report['generated_at']}")
    cell.font = Font(name="Arial", size=9, color="7A8F88")
    cell.alignment = Alignment(horizontal="center")

    for index, (label, col_width) in enumerate(columns, start=1):
        cell = ws.cell(row=header_row, column=index, value=label)
        cell.font = header_font
        cell.fill = accent_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws.column_dimensions[get_column_letter(index)].width = col_width
    ws.row_dimensions[header_row].height = 28
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(columns))}{ws.max_row}"


def daily_workbook(report: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "تقرير يومي"
    columns = [
        ("الرقم", 7),
        ("رقم التعريف", 22),
        ("اللقب", 18),
        ("الاسم", 18),
        ("القسم", 12),
        ("الفترة", 16),
        ("الوقت", 10),
        ("الحالة", 13),
        ("حالة SMS", 14),
        ("ملاحظة", 32),
    ]
    header_row = 6
    for index, label in enumerate([label for label, _ in columns], start=1):
        ws.cell(row=header_row, column=index, value=label)
    for index, item in enumerate(report["items"], start=1):
        row_index = header_row + index
        values = [
            index,
            item["student_nin"],
            item["last_name"],
            item["first_name"],
            item["class_group"],
            item["period_name"],
            item["scan_time"] or "—",
            item["status_label"],
            item["sms_status_label"] + (" (اختبار)" if item.get("sms_is_test") else ""),
            item["note"],
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col, value=value)
    _style_workbook(ws, report["title"], report, columns, header_row)

    # Summary strip after the data.
    summary_row = ws.max_row + 2
    ws.cell(row=summary_row, column=1, value="الملخص").font = Font(name="Arial", bold=True)
    labels = [("حاضر", "present"), ("متأخر", "late"), ("غائب", "absent"), ("مبرر", "excused")]
    col = 2
    for label, key in labels:
        cell = ws.cell(row=summary_row, column=col, value=f"{label}: {report['summary'].get(key, 0)}")
        cell.font = Font(name="Arial", bold=True, color="0B3B31")
        cell.fill = PatternFill("solid", fgColor="E7F3EE")
        col += 1
    ws.cell(row=summary_row, column=col, value=f"المجموع: {report['total']}").font = Font(name="Arial", bold=True)

    thin = Side(style="thin", color="DCE7E3")
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + len(report["items"]), min_col=1, max_col=len(columns)):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
            cell.border = Border(bottom=Side(style="hair", color="DCE7E3"))
        if (row[0].row - header_row) % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="F6FAF8")
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + len(report["items"])):
        note_cell = row[9] if len(row) > 9 else None
        if note_cell is not None:
            note_cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def monthly_workbook(report: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "تقرير شهري"
    columns = [
        ("الرقم", 7),
        ("رقم التعريف", 22),
        ("اللقب", 18),
        ("الاسم", 18),
        ("القسم", 12),
        ("حاضر", 10),
        ("متأخر", 10),
        ("غائب", 10),
        ("غياب مبرر", 12),
        ("مجموع السجلات", 14),
    ]
    header_row = 6
    for index, (label, _) in enumerate(columns, start=1):
        ws.cell(row=header_row, column=index, value=label)
    for index, item in enumerate(report["items"], start=1):
        row_index = header_row + index
        values = [
            index, item["student_nin"], item["last_name"], item["first_name"], item["class_group"],
            item["present"], item["late"], item["absent"], item["excused"], item["total"],
        ]
        for col, value in enumerate(values, start=1):
            ws.cell(row=row_index, column=col, value=value)
    _style_workbook(ws, report["title"], report, columns, header_row)
    totals_row = ws.max_row + 1
    totals = report["totals"]
    values = ["", "", "", "المجموع", "", totals["present"], totals["late"], totals["absent"], totals["excused"], totals["total"]]
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=totals_row, column=col, value=value)
        cell.font = Font(name="Arial", bold=True)
        cell.fill = PatternFill("solid", fgColor="E7F3EE")
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + len(report["items"])):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(bottom=Side(style="hair", color="DCE7E3"))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
