from __future__ import annotations

import hashlib
import json
import re
import secrets
import sqlite3
import unicodedata
from datetime import date, datetime, timedelta
from io import BytesIO
from pathlib import Path
from typing import Any

from bs4 import BeautifulSoup

from .utils import (
    ValidationError,
    clean_name,
    normalize_date,
    normalize_nin,
    normalize_phone,
    squash,
    validate_nin,
)

MAX_ROWS = 50_000
MAX_COLUMNS = 80
MAX_PHOTO_BYTES = 5 * 1024 * 1024
ALLOWED_PHOTO_EXT = {".jpg", ".jpeg", ".png", ".webp"}

FIELD_DEFS = [
    {
        "key": "student_nin",
        "label": "رقم التعريف الوطني",
        "hint": "الرقم التعريفي المستعمل في الباركود (إلزامي)",
        "required": True,
    },
    {"key": "last_name", "label": "اللقب", "hint": "لقب التلميذ", "required": False},
    {"key": "first_name", "label": "الاسم", "hint": "اسم التلميذ", "required": False},
    {"key": "birth_date", "label": "تاريخ الميلاد", "hint": "يوم/شهر/سنة أو Excel date", "required": False},
    {"key": "class_group", "label": "القسم / الفوج", "hint": "مثال: 1AS1", "required": False},
    {"key": "parent_phone", "label": "رقم هاتف الولي", "hint": "رقم جزائري يبدأ بـ 0", "required": False},
    {
        "key": "photo_path",
        "label": "مسار صورة التلميذ",
        "hint": "مسار ملف محلي على هذا الحاسوب (اختياري)",
        "required": False,
    },
]
FIELD_KEYS = [field["key"] for field in FIELD_DEFS]

KEYWORDS: dict[str, list[str]] = {
    "student_nin": [
        "رقم التعريف الوطني", "الرقم التعريفي", "رقم التعريف", "رقم تعريف",
        "التعريف الوطني", "رقم التسجيل", "رقم البطاقة", "رقم الملف",
        "identifiant", "num ident", "nni", "nin", "id",
    ],
    "last_name": ["اللقب العائلي", "اللقب", "لقب", "nom de famille", "nom", "last name"],
    "first_name": ["الاسم الشخصي", "الاسم", "الإسم", "اسم", "prénom", "prenom", "first name"],
    "birth_date": [
        "تاريخ الميلاد", "تاريخ ميلاد", "الميلاد", "date de naissance", "naissance",
        "birth date", "birthdate", "ddn",
    ],
    "class_group": [
        "القسم والفوج", "القسم / الفوج", "القسم/الفوج", "القسم", "الفوج",
        "classe", "groupe", "section", "class", "group",
    ],
    "parent_phone": [
        "رقم هاتف الولي", "هاتف الولي", "رقم الهاتف", "الهاتف", "هاتف",
        "رقم التليفون", "الموبايل", "الجوال", "telephone", "téléphone", "tel",
        "phone", "mobile", "gsm",
    ],
    "photo_path": ["مسار الصورة", "صورة التلميذ", "الصورة", "صورة", "photo", "image"],
}

def _normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKC", squash(value)).lower()
    text = re.sub(r"[\u064b-\u065f\u0670\u0640]", "", text)
    text = text.replace("أ", "ا").replace("إ", "ا").replace("آ", "ا")
    text = text.replace("ى", "ي").replace("ئ", "ي").replace("ؤ", "و")
    text = text.replace("ة", "ه")
    text = re.sub(r"[^\w\u0600-\u06ff]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


KEYWORDS_NORMALIZED = {
    field: [_normalize_header(keyword) for keyword in keywords]
    for field, keywords in KEYWORDS.items()
}


def _header_score(cells: list[str]) -> tuple[int, int]:
    normalized = [_normalize_header(cell) for cell in cells]
    joined = " | ".join(normalized)
    matched_fields = 0
    hits = 0
    for field, keywords in KEYWORDS_NORMALIZED.items():
        field_hit = False
        for keyword in keywords:
            if keyword and keyword in joined:
                field_hit = True
                hits += 2 if any(keyword == cell for cell in normalized) else 1
        if field_hit:
            matched_fields += 1
    return matched_fields, hits


def _decode_html(data: bytes) -> str:
    if data.startswith(b"\xff\xfe") or data.startswith(b"\xfe\xff"):
        try:
            return data.decode("utf-16")
        except UnicodeDecodeError:
            pass
    if data.startswith(b"\xef\xbb\xbf"):
        try:
            return data.decode("utf-8-sig")
        except UnicodeDecodeError:
            pass
    prefix = data[:8192].decode("latin-1", errors="ignore")
    meta = re.search(r"charset\s*=\s*[\"']?\s*([\w-]+)", prefix, re.IGNORECASE)
    candidates = []
    if meta:
        candidates.append(meta.group(1))
    candidates.extend(["utf-8", "cp1256", "windows-1252"])
    for encoding in dict.fromkeys(candidates):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def _sniff_kind(data: bytes) -> str:
    if not data:
        raise ValidationError("الملف فارغ.")
    # Definitive binary container signatures take precedence over text markers
    # embedded inside spreadsheet cells.
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    if data.startswith(b"PK\x03\x04"):
        return "xlsx"
    head = data[:65536]
    text_probe = (
        head.decode("utf-16", errors="ignore")
        if head[:2] in (b"\xff\xfe", b"\xfe\xff")
        else head.decode("latin-1", errors="ignore")
    )
    lowered = text_probe.lower()
    html_markers = ("<html", "<table", "<tr", "<!doctype html", "<head", "<body")
    if any(marker in lowered for marker in html_markers):
        return "html"
    raise ValidationError(
        "نوع الملف غير مدعوم. المدعوم: XLSX، XLS الثنائي، أو جدول HTML محفوظ بامتداد .xls."
    )


def _format_number(value: float | int) -> str:
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".15g")
    return str(value)


def _format_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, date):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, (int, float)):
        return _format_number(value)
    return squash(str(value))


def _clean_row(cells: list[str]) -> list[str]:
    row = [squash(cell) for cell in cells[:MAX_COLUMNS]]
    while row and not row[-1]:
        row.pop()
    return row


def _collect_html_rows(data: bytes) -> tuple[list[tuple[int, list[str]]], str]:
    html = _decode_html(data)
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "template"]):
        tag.decompose()
    tables = soup.find_all("table")
    if not tables:
        raise ValidationError("لم يتم العثور على أي جدول داخل ملف HTML.")

    best_rows: list[tuple[int, list[str]]] = []
    best_score = -1
    for table in tables:
        collected: list[tuple[int, list[str]]] = []
        for row_number, tr in enumerate(table.find_all("tr"), start=1):
            cells: list[str] = []
            for cell in tr.find_all(["th", "td"], recursive=False):
                text = squash(cell.get_text(" ", strip=True))
                cells.append(text)
                try:
                    colspan = int(cell.get("colspan", 1))
                except (TypeError, ValueError):
                    colspan = 1
                if colspan > 1:
                    cells.extend([""] * (min(colspan, 20) - 1))
            cleaned = _clean_row(cells)
            if cleaned:
                collected.append((row_number, cleaned))
        scored = sum(len(row) for _, row in collected)
        if scored > best_score:
            best_score = scored
            best_rows = collected
    if not best_rows:
        raise ValidationError("الجدول الموجود في الملف لا يحتوي على بيانات.")
    return best_rows, "HTML table"


def _collect_xlsx_rows(data: bytes) -> tuple[list[tuple[int, list[str]]], str]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        raise ValidationError(f"تعذر قراءة ملف Excel: {exc}") from exc
    if not workbook.worksheets:
        raise ValidationError("ملف Excel لا يحتوي على أوراق عمل.")
    sheet = max(
        workbook.worksheets,
        key=lambda ws: (ws.max_row or 0) * (ws.max_column or 0),
    )
    rows: list[tuple[int, list[str]]] = []
    try:
        for index, values in enumerate(sheet.iter_rows(values_only=True), start=1):
            cells = _clean_row([_format_cell(value) for value in values])
            if cells:
                rows.append((index, cells))
            if index > MAX_ROWS + 1000:
                raise ValidationError(f"الملف كبير جدًا (الحد {MAX_ROWS} صف).")
    finally:
        workbook.close()
    return rows, f"XLSX — {sheet.title}"


def _collect_xls_rows(data: bytes) -> tuple[list[tuple[int, list[str]]], str]:
    import xlrd

    try:
        workbook = xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise ValidationError(f"تعذر قراءة ملف XLS: {exc}") from exc
    if workbook.nsheets == 0:
        raise ValidationError("ملف XLS لا يحتوي على أوراق عمل.")
    sheet = max(workbook.sheets(), key=lambda sh: sh.nrows * sh.ncols)
    rows: list[tuple[int, list[str]]] = []
    for row_index in range(sheet.nrows):
        if row_index > MAX_ROWS + 1000:
            raise ValidationError(f"الملف كبير جدًا (الحد {MAX_ROWS} صف).")
        values: list[str] = []
        for col_index in range(sheet.ncols):
            cell = sheet.cell(row_index, col_index)
            if cell.ctype == xlrd.XL_CELL_DATE:
                try:
                    values.append(
                        xlrd.xldate_as_datetime(cell.value, workbook.datemode).strftime("%d/%m/%Y")
                    )
                except (ValueError, OverflowError):
                    values.append("")
            elif cell.ctype == xlrd.XL_CELL_NUMBER:
                values.append(_format_number(cell.value))
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                values.append("1" if cell.value else "0")
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                values.append("")
            else:
                values.append(squash(cell.value))
        cleaned = _clean_row(values)
        if cleaned:
            rows.append((row_index + 1, cleaned))
    return rows, f"XLS — {sheet.name}"


def _detect_header(rows: list[tuple[int, list[str]]]) -> int:
    best_index = 0
    best_score = (-1, -1)
    for index, (_, cells) in enumerate(rows[:20]):
        score = _header_score(cells)
        if score > best_score:
            best_score = score
            best_index = index
    return best_index


def _unique_columns(header: list[str], width: int) -> list[str]:
    columns: list[str] = []
    seen: dict[str, int] = {}
    for index in range(width):
        label = header[index] if index < len(header) else ""
        label = squash(label) or f"عمود {index + 1}"
        label = label[:80]
        count = seen.get(label, 0) + 1
        seen[label] = count
        columns.append(label if count == 1 else f"{label} ({count})")
    return columns


def suggest_mapping(columns: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {}
    used: set[int] = set()
    scored_headers = [_normalize_header(column) for column in columns]
    for field in FIELD_KEYS:
        best_column = None
        best_score = 0
        for index, header in enumerate(scored_headers):
            if index in used or not header:
                continue
            score = 0
            for keyword in KEYWORDS_NORMALIZED[field]:
                if not keyword:
                    continue
                if header == keyword:
                    score = max(score, 5)
                elif keyword in header:
                    score = max(score, 3)
            if score > best_score:
                best_score = score
                best_column = index
        if best_column is not None and best_score >= 3:
            mapping[field] = best_column
            used.add(best_column)
        else:
            mapping[field] = None
    return mapping


def _safe_token(token: str) -> str:
    if not re.fullmatch(r"[A-Za-z0-9_-]{16,80}", token or ""):
        raise ValidationError("جلسة استيراد غير صالحة. أعد رفع الملف.")
    return token


def _temp_path(tmp_dir: Path, token: str) -> Path:
    return Path(tmp_dir) / f"import_{_safe_token(token)}.json"


def _cleanup_tmp(tmp_dir: Path) -> None:
    now = datetime.now()
    for path in Path(tmp_dir).glob("import_*.json"):
        try:
            if now - datetime.fromtimestamp(path.stat().st_mtime) > timedelta(hours=3):
                path.unlink(missing_ok=True)
        except OSError:
            continue


def parse_file(tmp_dir: Path, filename: str, data: bytes) -> dict[str, Any]:
    if len(data) > 25 * 1024 * 1024:
        raise ValidationError("حجم الملف يتجاوز 25 ميغابايت.")
    Path(tmp_dir).mkdir(parents=True, exist_ok=True)
    _cleanup_tmp(tmp_dir)
    kind = _sniff_kind(data)
    if kind == "html":
        rows, source_type = _collect_html_rows(data)
    elif kind == "xlsx":
        rows, source_type = _collect_xlsx_rows(data)
    else:
        rows, source_type = _collect_xls_rows(data)

    if not rows:
        raise ValidationError("لم يتم العثور على صفوف قابلة للقراءة في الملف.")
    header_index = _detect_header(rows)
    header_row_number, raw_header = rows[header_index]
    width = min(MAX_COLUMNS, max(len(cells) for _, cells in rows))
    columns = _unique_columns(raw_header, width)
    data_rows = []
    for row_number, cells in rows[header_index + 1 :]:
        padded = (cells + [""] * width)[:width]
        if not any(padded):
            continue
        data_rows.append({"n": row_number, "cells": padded})
        if len(data_rows) > MAX_ROWS:
            raise ValidationError(f"الملف يحتوي على أكثر من {MAX_ROWS} صف.")
    if not data_rows:
        raise ValidationError("لم توجد صفوف بيانات بعد صف العناوين.")

    mapping = suggest_mapping(columns)
    token = secrets.token_urlsafe(24)
    payload = {
        "token": token,
        "filename": squash(filename)[:180] or "ملف بدون اسم",
        "source_type": source_type,
        "header_row_number": header_row_number,
        "columns": columns,
        "rows": data_rows,
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    _temp_path(tmp_dir, token).write_text(
        json.dumps(payload, ensure_ascii=False), encoding="utf-8"
    )
    return {
        "token": token,
        "file_name": payload["filename"],
        "file_type": source_type,
        "header_row": header_row_number,
        "columns": [{"index": index, "label": label} for index, label in enumerate(columns)],
        "suggested_mapping": mapping,
        "preview": [
            {"row": row["n"], "cells": row["cells"][:width]}
            for row in data_rows[:5]
        ],
        "total_rows": len(data_rows),
        "fields": FIELD_DEFS,
    }


def load_import(tmp_dir: Path, token: str) -> dict[str, Any]:
    path = _temp_path(tmp_dir, token)
    if not path.exists():
        raise ValidationError("انتهت صلاحية جلسة الاستيراد. أعد رفع الملف.")
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise ValidationError("تعذر قراءة جلسة الاستيراد. أعد رفع الملف.") from exc
    return payload


def delete_import(tmp_dir: Path, token: str) -> None:
    try:
        _temp_path(tmp_dir, token).unlink(missing_ok=True)
    except OSError:
        pass


def _validate_mapping(columns: list[str], mapping: dict[str, Any]) -> dict[str, int]:
    if not isinstance(mapping, dict):
        raise ValidationError("بيانات ربط الأعمدة غير صالحة.")
    clean: dict[str, int] = {}
    used_columns: dict[int, str] = {}
    for field in FIELD_KEYS:
        raw_index = mapping.get(field)
        if raw_index in (None, "", "null"):
            continue
        try:
            index = int(raw_index)
        except (TypeError, ValueError):
            raise ValidationError("بيانات ربط الأعمدة غير صالحة.")
        if index < 0 or index >= len(columns):
            raise ValidationError("أحد الأعمدة المختارة غير موجود في الملف.")
        if index in used_columns:
            raise ValidationError(
                f"العمود «{columns[index]}» مربوط بحقلين. كل حقل يحتاج عمودًا مستقلًا."
            )
        used_columns[index] = field
        clean[field] = index
    if "student_nin" not in clean:
        raise ValidationError("يجب ربط عمود «رقم التعريف الوطني» قبل المتابعة.")
    return clean


def _excel_serial_to_date(value: str) -> str | None:
    if not value.isdigit():
        return None
    serial = int(value)
    if 20000 <= serial <= 80000:
        try:
            return (date(1899, 12, 30) + timedelta(days=serial)).isoformat()
        except OverflowError:
            return None
    return None


def _build_records(
    payload: dict[str, Any], mapping: dict[str, Any]
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    columns = payload["columns"]
    clean_mapping = _validate_mapping(columns, mapping)
    records: dict[str, dict[str, Any]] = {}
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    stats = {
        "rows": 0,
        "valid": 0,
        "duplicates": 0,
        "invalid": 0,
        "no_phone": 0,
        "bad_phone": 0,
        "bad_birth": 0,
        "no_name": 0,
        "no_class": 0,
        "photos": 0,
    }

    def raw(field: str, cells: list[str]) -> str:
        index = clean_mapping.get(field)
        if index is None or index >= len(cells):
            return ""
        return squash(cells[index])

    for row in payload["rows"]:
        stats["rows"] += 1
        cells = row["cells"]
        row_number = row["n"]
        raw_nin = raw("student_nin", cells)
        if not raw_nin:
            allowed_other = clean_mapping.get("last_name") or clean_mapping.get("first_name")
            if allowed_other is not None:
                other = raw("last_name", cells) or raw("first_name", cells)
                if not other:
                    continue
            errors.append({"row": row_number, "message": "رقم التعريف الوطني فارغ."})
            stats["invalid"] += 1
            continue
        nin = normalize_nin(raw_nin)
        if re.search(r"\d[eE][+-]?\d+", raw_nin):
            errors.append({
                "row": row_number,
                "message": "رقم التعريف يبدو بصيغة علمية (Scientific notation). استعمل تنسيق نص في Excel.",
                "value": raw_nin,
            })
            stats["invalid"] += 1
            continue
        try:
            validate_nin(nin)
        except ValidationError as exc:
            errors.append({"row": row_number, "message": str(exc), "value": raw_nin})
            stats["invalid"] += 1
            continue

        first_name = clean_name(raw("first_name", cells))[:120]
        last_name = clean_name(raw("last_name", cells))[:120]
        class_group = clean_name(raw("class_group", cells))[:80]
        raw_birth = raw("birth_date", cells)
        birth_date = normalize_date(raw_birth) if raw_birth else None
        if raw_birth and not birth_date:
            birth_date = _excel_serial_to_date(raw_birth)
        if raw_birth and not birth_date:
            stats["bad_birth"] += 1
            warnings.append({
                "row": row_number,
                "message": f"تاريخ ميلاد غير مفهوم: «{raw_birth}» — سيُترك فارغًا.",
            })
        raw_phone = raw("parent_phone", cells)
        phone = normalize_phone(raw_phone) if raw_phone else ""
        if raw_phone and not re.fullmatch(r"\+?\d{8,15}", phone):
            stats["bad_phone"] += 1
            warnings.append({
                "row": row_number,
                "message": f"رقم الهاتف غير صالح: «{raw_phone}» — لن تُرسل رسائل SMS له.",
            })
            phone = ""
        if not phone:
            stats["no_phone"] += 1
        if not first_name and not last_name:
            stats["no_name"] += 1
            errors.append({"row": row_number, "message": "اسم التلميذ ولقبه فارغان معًا."})
            stats["invalid"] += 1
            continue
        if not class_group:
            stats["no_class"] += 1
        photo_path = raw("photo_path", cells)
        if photo_path:
            stats["photos"] += 1

        record = {
            "row": row_number,
            "student_nin": nin,
            "first_name": first_name,
            "last_name": last_name,
            "birth_date": birth_date,
            "class_group": class_group,
            "parent_phone": phone,
            "photo_path": photo_path,
            "_mapped": set(clean_mapping.keys()),
        }
        if nin in records:
            stats["duplicates"] += 1
            warnings.append({
                "row": row_number,
                "message": f"رقم التعريف {nin} مكرر في الملف — سيُعتمد آخر صف.",
            })
        records[nin] = record

    result = list(records.values())
    stats["valid"] = len(result)
    return result, errors, warnings, stats


def _existing_nins(conn: sqlite3.Connection, nins: list[str]) -> set[str]:
    existing: set[str] = set()
    for index in range(0, len(nins), 400):
        chunk = nins[index : index + 400]
        placeholders = ",".join("?" for _ in chunk)
        for row in conn.execute(
            f"SELECT student_nin FROM students WHERE student_nin IN ({placeholders})", chunk
        ).fetchall():
            existing.add(row["student_nin"].lower())
    return existing


def preview_import(
    conn: sqlite3.Connection, tmp_dir: Path, token: str, mapping: dict[str, Any]
) -> dict[str, Any]:
    payload = load_import(tmp_dir, token)
    records, errors, warnings, stats = _build_records(payload, mapping)
    existing = _existing_nins(conn, [record["student_nin"] for record in records])
    new_count = sum(1 for record in records if record["student_nin"].lower() not in existing)
    old_count = len(records) - new_count
    preview = [
        {
            "row": record["row"],
            "student_nin": record["student_nin"],
            "first_name": record["first_name"],
            "last_name": record["last_name"],
            "birth_date": record["birth_date"],
            "class_group": record["class_group"],
            "parent_phone": record["parent_phone"],
            "status": "تحديث" if record["student_nin"].lower() in existing else "جديد",
        }
        for record in records[:8]
    ]
    return {
        "summary": {
            **stats,
            "new": new_count,
            "existing": old_count,
            "invalid": len(errors),
        },
        "preview": preview,
        "errors": errors[:100],
        "warnings": warnings[:100],
        "errors_total": len(errors),
        "warnings_total": len(warnings),
    }


def _valid_phone(phone: str) -> bool:
    return bool(re.fullmatch(r"\+?\d{8,15}", phone or ""))


def _photo_magic_ok(data: bytes) -> tuple[bool, str]:
    if data.startswith(b"\xff\xd8\xff"):
        return True, ".jpg"
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return True, ".png"
    if data.startswith(b"RIFF") and data[8:12] == b"WEBP":
        return True, ".webp"
    return False, ""


def _copy_photo(source_value: str, photos_dir: Path) -> tuple[str | None, str | None]:
    value = squash(source_value)
    if not value:
        return None, None
    if value.lower().startswith(("http://", "https://")):
        return None, "روابط الصور غير مدعومة. ارفع الصورة من صفحة التلاميذ."
    source = Path(value).expanduser()
    if not source.is_absolute():
        return None, "مسار الصورة يجب أن يكون مسارًا كاملًا على هذا الحاسوب، أو ارفع الصورة يدويًا."
    if not source.exists() or not source.is_file():
        return None, "ملف الصورة غير موجود في المسار المحدد."
    if source.suffix.lower() not in ALLOWED_PHOTO_EXT:
        return None, "صيغة الصورة غير مدعومة (JPG، PNG، WEBP فقط)."
    try:
        if source.stat().st_size > MAX_PHOTO_BYTES:
            return None, "حجم الصورة يتجاوز 5 ميغابايت."
        data = source.read_bytes()
    except OSError:
        return None, "تعذر قراءة ملف الصورة."
    ok, detected_ext = _photo_magic_ok(data)
    if not ok:
        return None, "الملف لا يبدو صورة صالحة."
    extension = detected_ext
    digest = hashlib.sha256(data).hexdigest()[:16]
    target_name = f"student_{digest}{extension}"
    photos_dir = Path(photos_dir)
    photos_dir.mkdir(parents=True, exist_ok=True)
    target = photos_dir / target_name
    if not target.exists():
        target.write_bytes(data)
    return f"/media/photos/{target_name}", None


def commit_import(
    conn: sqlite3.Connection,
    tmp_dir: Path,
    photos_dir: Path,
    token: str,
    mapping: dict[str, Any],
) -> dict[str, Any]:
    payload = load_import(tmp_dir, token)
    records, errors, warnings, stats = _build_records(payload, mapping)
    if not records:
        raise ValidationError("لا يوجد أي صف صالح للاستيراد. راجع الأخطاء المعروضة.")

    existing = _existing_nins(conn, [record["student_nin"] for record in records])
    inserted = 0
    updated = 0
    photos_copied = 0
    photo_failures: list[dict[str, Any]] = []
    try:
        with conn:
            for record in records:
                photo_url, photo_error = (None, None)
                if record["photo_path"]:
                    photo_url, photo_error = _copy_photo(record["photo_path"], photos_dir)
                    if photo_url:
                        photos_copied += 1
                    elif photo_error:
                        photo_failures.append(
                            {"row": record["row"], "message": photo_error}
                        )
                is_existing = record["student_nin"].lower() in existing
                mapped = record["_mapped"]
                if is_existing:
                    updates: list[str] = []
                    values: list[Any] = []
                    if "last_name" in mapped and record["last_name"]:
                        updates.append("last_name = ?")
                        values.append(record["last_name"])
                    if "first_name" in mapped and record["first_name"]:
                        updates.append("first_name = ?")
                        values.append(record["first_name"])
                    if "birth_date" in mapped and record["birth_date"]:
                        updates.append("birth_date = ?")
                        values.append(record["birth_date"])
                    if "class_group" in mapped and record["class_group"]:
                        updates.append("class_group = ?")
                        values.append(record["class_group"])
                    if "parent_phone" in mapped and _valid_phone(record["parent_phone"]):
                        updates.append("parent_phone = ?")
                        values.append(record["parent_phone"])
                    if photo_url:
                        updates.append("photo_url = ?")
                        values.append(photo_url)
                    updates.extend([
                        "is_active = 1",
                        "updated_at = datetime('now', 'localtime')",
                    ])
                    conn.execute(
                        f"UPDATE students SET {', '.join(updates)} WHERE student_nin = ?",
                        [*values, record["student_nin"]],
                    )
                    updated += 1
                else:
                    conn.execute(
                        """INSERT INTO students
                           (student_nin, first_name, last_name, birth_date, class_group,
                            parent_phone, photo_url, is_active)
                           VALUES (?, ?, ?, ?, ?, ?, ?, 1)""",
                        (
                            record["student_nin"],
                            record["first_name"],
                            record["last_name"],
                            record["birth_date"],
                            record["class_group"],
                            record["parent_phone"],
                            photo_url,
                        ),
                    )
                    inserted += 1
    except sqlite3.Error as exc:
        raise ValidationError(f"فشل حفظ البيانات في قاعدة البيانات: {exc}") from exc

    delete_import(tmp_dir, token)
    result_warnings = warnings[:100]
    result_warnings.extend(
        {"row": failure["row"], "message": failure["message"]}
        for failure in photo_failures[:30]
    )
    return {
        "inserted": inserted,
        "updated": updated,
        "total": inserted + updated,
        "errors": errors[:100],
        "errors_total": len(errors),
        "warnings": result_warnings,
        "warnings_total": len(warnings) + len(photo_failures),
        "photos_copied": photos_copied,
        "photo_failures": photo_failures[:30],
        "summary": {
            **stats,
            "new": inserted,
            "existing": updated,
            "invalid": len(errors),
        },
    }
