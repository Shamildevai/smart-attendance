from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook


def _seed_attendance(db, students, period_id: int):
    today = date.today().isoformat()
    statuses = ["PRESENT", "LATE", "ABSENT"]
    with db:
        for student, status in zip(students, statuses):
            db.execute(
                """INSERT INTO attendance_logs
                   (student_nin, period_id, attendance_date, scan_time, status, source)
                   VALUES (?, ?, ?, ?, ?, 'GATE')""",
                (
                    student["student_nin"],
                    period_id,
                    today,
                    "07:50:00" if status == "PRESENT" else ("08:20:00" if status == "LATE" else None),
                    status,
                ),
            )


def test_daily_and_monthly_reports(client, db, make_student, period):
    students = [
        make_student(first_name="أمينة", last_name="بوقرة", class_group="1AS3"),
        make_student(first_name="إلياس", last_name="مرزوقي", class_group="1AS3"),
        make_student(first_name="نسرين", last_name="حديد", class_group="1AS3"),
    ]
    _seed_attendance(db, students, period["id"])
    today = date.today().isoformat()

    daily = client.get(f"/api/reports/daily?date={today}&period_id={period['id']}")
    assert daily.status_code == 200
    report = daily.get_json()
    assert report["summary"] == {"PRESENT": 1, "LATE": 1, "ABSENT": 1, "EXCUSED": 0}
    assert report["total"] == 3

    workbook_response = client.get(f"/api/reports/daily.xlsx?date={today}")
    assert workbook_response.status_code == 200
    assert workbook_response.data[:2] == b"PK"
    workbook = load_workbook(BytesIO(workbook_response.data))
    assert "تقرير يومي" in workbook.sheetnames
    assert workbook["تقرير يومي"].max_row >= 4

    month = today[:7]
    monthly = client.get(f"/api/reports/monthly?month={month}&period_id={period['id']}")
    assert monthly.status_code == 200
    assert monthly.get_json()["ok"] is True

    monthly_file = client.get(f"/api/reports/monthly.xlsx?month={month}")
    assert monthly_file.status_code == 200
    book = load_workbook(BytesIO(monthly_file.data))
    assert "تقرير شهري" in book.sheetnames

    printable = client.get(f"/reports/print?date={today}&period_id={period['id']}")
    assert printable.status_code == 200
    assert "بوقرة" in printable.get_data(as_text=True)


def test_barcode_qr_and_print_sheets(client, make_student, period):
    student = make_student(student_nin="2099000401", first_name="ريان", last_name="بلحاج", class_group="1AS4")
    nin = student["student_nin"]

    barcode = client.get(f"/api/barcode/{nin}")
    assert barcode.status_code == 200
    assert barcode.mimetype == "image/svg+xml"
    assert "<svg" in barcode.get_data(as_text=True)

    qr = client.get(f"/api/qr/{nin}")
    assert qr.status_code == 200
    assert qr.mimetype == "image/svg+xml"
    assert "<svg" in qr.get_data(as_text=True)

    cards = client.get("/print/sheet?mode=cards&scope=all&layout=10")
    assert cards.status_code == 200
    html = cards.get_data(as_text=True)
    assert "sheet--cards-10" in html
    assert "ريان" in html
    assert html.count("print-card") >= 1

    stickers = client.get("/print/sheet?mode=stickers&scope=all&layout=30")
    assert stickers.status_code == 200
    sticker_html = stickers.get_data(as_text=True)
    assert "sheet--stickers-30" in sticker_html
    assert "print-sticker" in sticker_html

    by_class = client.get("/print/sheet?mode=cards&scope=class&class_group=1AS4&layout=8")
    assert by_class.status_code == 200
    assert "sheet--cards-8" in by_class.get_data(as_text=True)


def test_local_backup_download(client, make_student):
    make_student(student_nin="2099000411")
    response = client.get("/api/system/backup")
    assert response.status_code == 200
    assert response.data.startswith(b"SQLite format 3\x00")
    assert ".db" in response.headers.get("Content-Disposition", "")
